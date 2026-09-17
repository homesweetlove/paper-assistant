# -*- coding: utf-8 -*-
"""
논문 작성 보조 도구 (오프라인)
- .docx 파일 불러오기 / 저장
- 통계: 글자수, 단어수, 문장수, 문단수, 평균 문장 길이
- 문장 검사: 과도하게 긴 문장, 감탄/의문문, 조사 중복, 어미 반복
- 문체 검사: 구어체 종결어미, 문두 접속어, 1인칭, 수식어 남용
- 반복 어휘: 단어 빈도 분석
실행: python paper_assistant.py
"""
import re
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    from docx import Document
except ImportError:
    Document = None

SENT_SPLIT = re.compile(r"[^.!?\n]+(?:[.!?]+|$)")
TOKEN_RE = re.compile(r"[가-힣a-zA-Z0-9]+")
REPEAT_TOKEN_RE = re.compile(r"\b([가-힣]+)\s+\1\b")

COLLOQUIAL_ENDINGS = re.compile(
    r"(해요|했어요|이에요|예요|이예요|죠|네요|거든요|잖아요|군요|구나|어라|지요)\s*[.!?]?$"
)
SENT_OPENERS = ("그래서", "근데", "그리고", "하지만", "그런데", "그러면", "또")
INTENSIFIERS = ("매우", "정말", "아주", "굉장히", "엄청", "되게", "진짜", "너무")
FIRST_PERSON = ("내가", "저는", "우리는", "우리가")
PARTICLES = ("은", "는", "이", "가", "을", "를", "의", "에", "에서", "로", "으로", "와", "과", "도", "만")

ISSUE_SENT = "문장"
ISSUE_STYLE = "문체"
ISSUE_REPEAT = "반복"


def split_sentences(text):
    """(문장, 시작오프셋, 끝오프셋) 리스트 반환."""
    out = []
    for m in SENT_SPLIT.finditer(text):
        s = m.group(0)
        if s.strip():
            start = m.start() + (len(s) - len(s.lstrip()))
            out.append((s.strip(), start, m.end()))
    return out


def analyze(text, long_sent_limit=100, repeat_ratio=0.02, repeat_min=4):
    issues = []  # (kind, msg, start, end)
    sentences = split_sentences(text)
    ending_streak = []  # 어미 반복 추적

    for sent, start, end in sentences:
        if len(sent) >= long_sent_limit:
            issues.append((ISSUE_SENT, f"문장이 깁니다 ({len(sent)}자): {sent[:30]}…", start, end))
        if sent.rstrip().endswith(("!", "?")):
            issues.append((ISSUE_SENT, f"감탄/의문문: {sent[:30]}…", start, end))
        m = COLLOQUIAL_ENDINGS.search(sent)
        if m:
            issues.append((ISSUE_STYLE, f"구어체 어미 '~{m.group(1)}': {sent[:30]}…", start, end))
        first = sent.split()[0] if sent.split() else ""
        if first.rstrip(",") in SENT_OPENERS:
            issues.append((ISSUE_STYLE, f"문두 접속어 '{first.rstrip(',')}': {sent[:30]}…", start, end))
        for w in FIRST_PERSON:
            if w in sent:
                issues.append((ISSUE_STYLE, f"1인칭 '{w}' — '본 연구에서는' 등 권장: {sent[:30]}…", start, end))
                break
        for w in INTENSIFIERS:
            if w in sent:
                issues.append((ISSUE_STYLE, f"수식어 '{w}' — 정량적 표현 권장: {sent[:30]}…", start, end))
                break
        for m2 in REPEAT_TOKEN_RE.finditer(sent):
            if m2.group(1) in PARTICLES or len(m2.group(1)) >= 2:
                issues.append((ISSUE_REPEAT, f"같은 단어 연속 '{m2.group(1)} {m2.group(1)}'", start + m2.start(), start + m2.end()))

        # 종결어미 추적 (~다. 로 끝나는 명사형 패턴의 꼬리)
        tail = re.search(r"([가-힣]{1,4})[.!?]?$", sent)
        tail = tail.group(1) if tail else ""
        if ending_streak and ending_streak[-1][0] == tail and tail:
            ending_streak.append((tail, start, end))
            if len(ending_streak) >= 3:
                issues.append((ISSUE_STYLE, f"'~{tail}' 어미가 {len(ending_streak)}문장 연속 반복", start, end))
        else:
            ending_streak = [(tail, start, end)]

    # 단어 빈도
    words = [w.lower() for w in TOKEN_RE.findall(text) if len(w) >= 2 and w not in PARTICLES]
    freq = {}
    for w in words:
        freq[w] = freq.get(w, 0) + 1
    total = len(words) or 1
    overused = sorted(
        ((w, c) for w, c in freq.items() if c >= repeat_min and c / total >= repeat_ratio),
        key=lambda x: -x[1],
    )
    for w, c in overused:
        issues.append((ISSUE_REPEAT, f"단어 '{w}' {c}회 사용 ({c / total:.1%}) — 과다 반복 가능", 0, 0))

    stats = {
        "글자수(공백포함)": len(text),
        "글자수(공백제외)": len(re.sub(r"\s", "", text)),
        "단어수": len(text.split()),
        "문장수": len(sentences),
        "문단수": len([p for p in text.split("\n") if p.strip()]),
        "평균 문장길이": round(sum(len(s) for s, _, _ in sentences) / max(len(sentences), 1), 1),
    }
    top_words = sorted(freq.items(), key=lambda x: -x[1])[:20]
    return issues, stats, top_words


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("논문 작성 보조 도구")
        self.geometry("1100x700")
        self.issues = []
        self._build_ui()

    def _build_ui(self):
        bar = ttk.Frame(self)
        bar.pack(fill="x", padx=6, pady=4)
        ttk.Button(bar, text="docx 열기", command=self.open_docx).pack(side="left")
        ttk.Button(bar, text="docx로 저장", command=self.save_docx).pack(side="left", padx=4)
        ttk.Button(bar, text="분석 실행", command=self.run_analysis).pack(side="left", padx=4)
        ttk.Button(bar, text="전체 복사", command=self.copy_all).pack(side="left", padx=4)
        ttk.Label(bar, text="긴 문장 기준:").pack(side="left", padx=(16, 2))
        self.limit_var = tk.IntVar(value=100)
        ttk.Spinbox(bar, from_=40, to=300, width=5, textvariable=self.limit_var).pack(side="left")
        ttk.Label(bar, text="자").pack(side="left")
        self.realtime_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(bar, text="실시간 분석", variable=self.realtime_var).pack(side="left", padx=(16, 0))

        paned = ttk.PanedWindow(self, orient="horizontal")
        paned.pack(fill="both", expand=True, padx=6, pady=4)

        left = ttk.Frame(paned)
        self.editor = tk.Text(left, wrap="word", font=("맑은 고딕", 11), undo=True)
        scroll = ttk.Scrollbar(left, command=self.editor.yview)
        self.editor.configure(yscrollcommand=scroll.set)
        self.editor.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        self.editor.tag_configure("issue", background="#ffe08a")
        self.editor.bind("<KeyRelease>", self._on_key)
        self._pending = None
        paned.add(left, weight=3)

        right = ttk.Frame(paned)
        nb = ttk.Notebook(right)
        nb.pack(fill="both", expand=True)

        self.stats_text = tk.Text(nb, wrap="word", font=("맑은 고딕", 10), state="disabled")
        nb.add(self.stats_text, text="통계")

        self.issue_tree = ttk.Treeview(nb, columns=("kind", "msg"), show="headings")
        self.issue_tree.heading("kind", text="유형")
        self.issue_tree.heading("msg", text="내용")
        self.issue_tree.column("kind", width=50)
        self.issue_tree.column("msg", width=380)
        self.issue_tree.bind("<<TreeviewSelect>>", self.jump_to_issue)
        nb.add(self.issue_tree, text="검사 결과")

        self.words_text = tk.Text(nb, wrap="word", font=("맑은 고딕", 10), state="disabled")
        nb.add(self.words_text, text="단어 빈도")

        paned.add(right, weight=2)

        self.status = ttk.Label(self, text="docx 파일을 열거나 텍스트를 직접 입력한 뒤 '분석 실행'을 누르세요.")
        self.status.pack(fill="x", padx=6, pady=2)

    def open_docx(self):
        if Document is None:
            messagebox.showerror("오류", "python-docx가 없습니다: pip install python-docx")
            return
        path = filedialog.askopenfilename(filetypes=[("Word 문서", "*.docx")])
        if not path:
            return
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.editor.delete("1.0", "end")
        self.editor.insert("1.0", text)
        self.status.config(text=f"불러옴: {path}")

    def save_docx(self):
        if Document is None:
            messagebox.showerror("오류", "python-docx가 없습니다: pip install python-docx")
            return
        path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word 문서", "*.docx")])
        if not path:
            return
        doc = Document()
        for para in self.editor.get("1.0", "end-1c").split("\n"):
            doc.add_paragraph(para)
        doc.save(path)
        self.status.config(text=f"저장됨: {path}")

    def copy_all(self):
        self.clipboard_clear()
        self.clipboard_append(self.editor.get("1.0", "end-1c"))
        self.status.config(text="전체 텍스트를 클립보드에 복사했습니다.")

    def _on_key(self, _event):
        if not self.realtime_var.get():
            return
        if self._pending:
            self.after_cancel(self._pending)
        self._pending = self.after(500, self.run_analysis)

    def run_analysis(self):
        self._pending = None
        text = self.editor.get("1.0", "end-1c")
        if not text.strip():
            self.status.config(text="분석할 텍스트가 없습니다.")
            return
        self.issues, stats, top_words = analyze(text, long_sent_limit=self.limit_var.get())

        self.editor.tag_remove("issue", "1.0", "end")
        for row in self.issue_tree.get_children():
            self.issue_tree.delete(row)
        for i, (kind, msg, s, e) in enumerate(self.issues):
            self.issue_tree.insert("", "end", iid=str(i), values=(kind, msg))
            if e > s:
                self.editor.tag_add("issue", f"1.0+{s}c", f"1.0+{e}c")

        self.stats_text.config(state="normal")
        self.stats_text.delete("1.0", "end")
        self.stats_text.insert("end", "\n".join(f"{k}: {v}" for k, v in stats.items()))
        self.stats_text.insert("end", f"\n\n검출된 이슈: {len(self.issues)}건")
        self.stats_text.config(state="disabled")

        self.words_text.config(state="normal")
        self.words_text.delete("1.0", "end")
        self.words_text.insert("end", "단어 빈도 TOP 20\n\n" + "\n".join(f"{w}: {c}회" for w, c in top_words))
        self.words_text.config(state="disabled")

        self.status.config(
            text=f"글자수 {stats['글자수(공백제외)']:,}자 | 단어 {stats['단어수']:,}개 | 문장 {stats['문장수']}개 | 이슈 {len(self.issues)}건"
        )

    def jump_to_issue(self, _event):
        sel = self.issue_tree.selection()
        if not sel:
            return
        _, _, s, e = self.issues[int(sel[0])]
        if e <= s:
            return
        self.editor.see(f"1.0+{s}c")
        self.editor.tag_remove("sel", "1.0", "end")
        self.editor.tag_add("sel", f"1.0+{s}c", f"1.0+{e}c")
        self.editor.focus_set()


if __name__ == "__main__":
    App().mainloop()
